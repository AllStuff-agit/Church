#!/usr/bin/env python3
"""Replace a worship PPT sermon section with grass-theme title and verse slides.

This script is intentionally repo-local and tuned for the church worship PPT
assets under /mnt/c/Oddments_Repository/교회/주일 예배자료.
"""

from __future__ import annotations

import argparse
import copy
import importlib.util
import os
import re
import shutil
import subprocess
import sys
import tempfile
import uuid
import zipfile
from dataclasses import dataclass
from pathlib import Path
from typing import Iterable

from lxml import etree
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

PML_NS = "http://schemas.openxmlformats.org/presentationml/2006/main"
A_NS = "http://schemas.openxmlformats.org/drawingml/2006/main"
P14_NS = "http://schemas.microsoft.com/office/powerpoint/2010/main"
XML_NS = "http://www.w3.org/XML/1998/namespace"
P14_EXT_URI = "{521415D9-36F7-43E2-AB2F-B90AF26B5E84}"
NS = {"p": PML_NS, "p14": P14_NS}

DEFAULT_BIBLE_ROOT = "/mnt/c/Oddments_Repository/교회/주일 예배자료/2. 개역개정성경ppt"
DEFAULT_SECTION_REPLACER = "/home/wl/workspace/playground/sandbox/.codex/skills/church-worship-ppt/scripts/ppt_section_replacer.py"
DEFAULT_ANIMATION_TEMPLATE = "/mnt/c/Oddments_Repository/교회/주일 예배자료/1. 주일오전예배ppt/template.pptx"

EMU_PER_INCH = 914400
VERSE_HEADER_NAMES = ("TextBox 1", "TextBox 8")
VERSE_BODY_NAMES = ("TextBox 2", "TextBox 15")
UNUSED_VERSE_NAMES = ("TextBox 22",)
TITLE_DESIGN_BASE_SIZE_PT = 48.0

BOOKS = [
    "창세기", "출애굽기", "레위기", "민수기", "신명기", "여호수아", "사사기", "룻기",
    "사무엘상", "사무엘하", "열왕기상", "열왕기하", "역대상", "역대하", "에스라", "느헤미야",
    "에스더", "욥기", "시편", "잠언", "전도서", "아가", "이사야", "예레미야", "예레미야애가",
    "에스겔", "다니엘", "호세아", "요엘", "아모스", "오바댜", "요나", "미가", "나훔", "하박국",
    "스바냐", "학개", "스가랴", "말라기", "마태복음", "마가복음", "누가복음", "요한복음",
    "사도행전", "로마서", "고린도전서", "고린도후서", "갈라디아서", "에베소서", "빌립보서",
    "골로새서", "데살로니가전서", "데살로니가후서", "디모데전서", "디모데후서", "디도서",
    "빌레몬서", "히브리서", "야고보서", "베드로전서", "베드로후서", "요한일서", "요한이서",
    "요한삼서", "유다서", "요한계시록",
]

ALIASES = {
    "창": "창세기", "출": "출애굽기", "레": "레위기", "민": "민수기", "신": "신명기",
    "수": "여호수아", "삿": "사사기", "룻": "룻기",
    "삼상": "사무엘상", "삼하": "사무엘하", "왕상": "열왕기상", "왕하": "열왕기하",
    "대상": "역대상", "대하": "역대하", "스": "에스라", "느": "느헤미야", "에": "에스더",
    "욥": "욥기", "시": "시편", "잠": "잠언", "전": "전도서", "아": "아가",
    "사": "이사야", "렘": "예레미야", "애": "예레미야애가", "겔": "에스겔", "단": "다니엘",
    "호": "호세아", "욜": "요엘", "암": "아모스", "옵": "오바댜", "욘": "요나", "미": "미가",
    "나": "나훔", "합": "하박국", "습": "스바냐", "학": "학개", "슥": "스가랴", "말": "말라기",
    "마": "마태복음", "막": "마가복음", "눅": "누가복음", "요": "요한복음", "행": "사도행전",
    "마태": "마태복음", "마가": "마가복음", "누가": "누가복음", "요한": "요한복음", "사도": "사도행전",
    "롬": "로마서", "고전": "고린도전서", "고후": "고린도후서", "갈": "갈라디아서", "엡": "에베소서",
    "빌": "빌립보서", "골": "골로새서", "살전": "데살로니가전서", "살후": "데살로니가후서",
    "딤전": "디모데전서", "딤후": "디모데후서", "딛": "디도서", "몬": "빌레몬서", "히": "히브리서",
    "약": "야고보서", "벧전": "베드로전서", "벧후": "베드로후서", "요일": "요한일서", "요이": "요한이서",
    "요삼": "요한삼서", "유": "유다서", "계": "요한계시록", "계시록": "요한계시록",
}

for b in BOOKS:
    ALIASES[b] = b


def book_alias_candidates() -> list[tuple[str, str]]:
    return sorted(ALIASES.items(), key=lambda kv: len(re.sub(r"\s+", "", kv[0])), reverse=True)


def starts_with_book_alias(text: str) -> bool:
    text_lower = re.sub(r"^\s+", "", text).lower()
    return any(text_lower.startswith(re.sub(r"\s+", "", alias).lower()) for alias, _ in book_alias_candidates())


def split_reference_parts(text: str) -> list[str]:
    parts: list[str] = []
    start = 0
    for i, ch in enumerate(text):
        if ch in "/;；" or (ch == "," and starts_with_book_alias(text[i + 1:])):
            part = text[start:i].strip()
            if part:
                parts.append(part)
            start = i + 1
    last = text[start:].strip()
    if last:
        parts.append(last)
    return parts


def normalize(s: str) -> str:
    return re.sub(r"[\s_\-·.()\[\]{}]", "", s).lower()


def as_path(path: str) -> Path:
    """Accept WSL paths and simple Windows C:\\ paths."""
    p = path.strip().strip('"')
    m = re.match(r"^([A-Za-z]):[\\/](.*)$", p)
    if m:
        drive = m.group(1).lower()
        rest = m.group(2).replace("\\", "/")
        return Path(f"/mnt/{drive}/{rest}")
    return Path(p)


@dataclass(frozen=True)
class RefRange:
    book: str
    start_chapter: int
    start_verse: int
    end_chapter: int
    end_verse: int
    # Each segment is (start_chapter, start_verse, end_chapter, end_verse).
    # This supports discontinuous readings like 사도행전 2:1-4, 14-18.
    segments: tuple[tuple[int, int, int, int], ...] = ()
    # Optional book per segment. Empty means every segment uses ``book``.
    # This supports multi-book readings like 마태복음 16:18/에베소서 2:20-22.
    segment_books: tuple[str, ...] = ()

    def iter_segments(self) -> tuple[tuple[int, int, int, int], ...]:
        return self.segments or ((self.start_chapter, self.start_verse, self.end_chapter, self.end_verse),)

    def iter_book_segments(self) -> tuple[tuple[str, int, int, int, int], ...]:
        segments = self.iter_segments()
        books = self.segment_books or tuple(self.book for _ in segments)
        if len(books) != len(segments):
            raise ValueError("본문 segment_books 개수가 segments 개수와 다릅니다")
        return tuple(
            (book, sc, sv, ec, ev)
            for book, (sc, sv, ec, ev) in zip(books, segments)
        )


@dataclass(frozen=True)
class VerseRef:
    book: str
    chapter: int
    verse: int


def _parse_single_book_reference(text: str, original: str) -> tuple[str, tuple[tuple[int, int, int, int], ...]]:
    # Keep ':', '-', and ',' because they carry chapter/verse range meaning.
    compact = text.strip().replace("：", ":").replace("~", "-").replace("–", "-").replace("—", "-").replace("，", ",")
    compact = re.sub(r"\s+", "", compact)
    candidates = book_alias_candidates()
    book = None
    rest = None
    compact_lower = compact.lower()
    for alias, canonical in candidates:
        alias_key = re.sub(r"\s+", "", alias).lower()
        if compact_lower.startswith(alias_key):
            book = canonical
            rest = compact[len(alias_key):]
            break
    if not book or not rest:
        raise ValueError(f"본문에서 성경 책 이름을 찾을 수 없습니다: {original}")

    # Support: 4:1-10, 4장1-10절, 4:1-5:3, 4장1절-5장3절,
    # and discontinuous same-chapter ranges like 2:1-4,14-18.
    rest = rest.replace("장", ":").replace("절", "")
    rest = rest.replace("~", "-").replace("–", "-").replace("—", "-").replace("，", ",")
    parts = [part for part in rest.split(",") if part]
    if not parts:
        raise ValueError(f"본문 범위를 해석할 수 없습니다: {original}")

    segments: list[tuple[int, int, int, int]] = []
    current_chapter: int | None = None
    for i, part in enumerate(parts):
        if i == 0 or ":" in part:
            m = re.fullmatch(r"(\d+)[:](\d+)(?:-(?:(\d+)[:])?(\d+))?", part)
            if not m:
                raise ValueError(f"본문 범위를 해석할 수 없습니다: {original}")
            sc = int(m.group(1))
            sv = int(m.group(2))
            ec = int(m.group(3)) if m.group(3) else sc
            ev = int(m.group(4)) if m.group(4) else sv
        else:
            if current_chapter is None:
                raise ValueError(f"쉼표 뒤 범위에 장 정보가 없습니다: {original}")
            m = re.fullmatch(r"(\d+)(?:-(?:(\d+)[:])?(\d+))?", part)
            if not m:
                raise ValueError(f"본문 범위를 해석할 수 없습니다: {original}")
            sc = current_chapter
            sv = int(m.group(1))
            ec = int(m.group(2)) if m.group(2) else sc
            ev = int(m.group(3)) if m.group(3) else sv
        if (ec, ev) < (sc, sv):
            raise ValueError(f"본문 종료 범위가 시작보다 앞섭니다: {original}")
        segments.append((sc, sv, ec, ev))
        current_chapter = ec

    return book, tuple(segments)


def parse_reference(text: str) -> RefRange:
    raw = text.strip()
    compact = raw.replace("：", ":").replace("~", "-").replace("–", "-").replace("—", "-").replace("，", ",")
    parts = split_reference_parts(compact)
    if not parts:
        raise ValueError(f"본문 범위를 해석할 수 없습니다: {text}")

    all_segments: list[tuple[int, int, int, int]] = []
    all_books: list[str] = []
    for part in parts:
        book, segments = _parse_single_book_reference(part, raw)
        all_segments.extend(segments)
        all_books.extend([book] * len(segments))

    if not all_segments:
        raise ValueError(f"본문 범위를 해석할 수 없습니다: {text}")
    first = all_segments[0]
    last = all_segments[-1]
    return RefRange(all_books[0], first[0], first[1], last[2], last[3], tuple(all_segments), tuple(all_books))


def _verse_part(sv: int, ev: int) -> str:
    return str(sv) if sv == ev else f"{sv}-{ev}"


def display_reference(r: RefRange) -> str:
    book_segments = r.iter_book_segments()
    books = {book for book, _, _, _, _ in book_segments}
    if len(books) > 1:
        parts = []
        for book, sc, sv, ec, ev in book_segments:
            if sc == ec:
                parts.append(f"{book} {sc}장 {_verse_part(sv, ev)}절")
            else:
                parts.append(f"{book} {sc}장 {sv}절-{ec}장 {ev}절")
        return ", ".join(parts)

    book = book_segments[0][0]
    segments = tuple((sc, sv, ec, ev) for _, sc, sv, ec, ev in book_segments)
    chapters = {sc for sc, _, ec, _ in segments} | {ec for _, _, ec, _ in segments}
    if len(chapters) == 1:
        chapter = segments[0][0]
        parts = [_verse_part(sv, ev) if sc == ec else f"{sc}장 {sv}절-{ec}장 {ev}절" for sc, sv, ec, ev in segments]
        return f"{book} {chapter}장 {', '.join(parts)}절"
    parts = []
    for sc, sv, ec, ev in segments:
        if sc == ec:
            parts.append(f"{sc}장 {_verse_part(sv, ev)}절")
        else:
            parts.append(f"{sc}장 {sv}절-{ec}장 {ev}절")
    return f"{book} {', '.join(parts)}"

def load_section_replacer(path: Path):
    if not path.exists():
        raise FileNotFoundError(f"church-worship-ppt 섹션 교체 스크립트 없음: {path}")
    spec = importlib.util.spec_from_file_location("ppt_section_replacer", path)
    if spec is None or spec.loader is None:
        raise RuntimeError(f"스크립트 로드 실패: {path}")
    mod = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(mod)
    return mod


def convert_to_pptx(src: Path, out_dir: Path) -> Path:
    if src.suffix.lower() == ".pptx":
        return src
    result = subprocess.run(
        ["libreoffice", "--headless", "--convert-to", "pptx", "--outdir", str(out_dir), str(src)],
        text=True,
        capture_output=True,
    )
    converted = out_dir / f"{src.stem}.pptx"
    if result.returncode != 0 or not converted.exists():
        raise RuntimeError(f"PPT 변환 실패: {src}\nSTDOUT:{result.stdout}\nSTDERR:{result.stderr}")
    return converted


def find_chapter_ppt(bible_root: Path, book: str, chapter: int) -> Path:
    if not bible_root.exists():
        raise FileNotFoundError(f"개역개정성경ppt 폴더 없음: {bible_root}")
    ch = str(chapter)
    ch2 = f"{chapter:02d}"
    exact_names = {normalize(f"{book}{ch}장"), normalize(f"{book}{ch2}장")}
    candidates: list[tuple[int, Path]] = []
    for f in bible_root.rglob("*"):
        if f.suffix.lower() not in (".ppt", ".pptx"):
            continue
        stem = normalize(f.stem)
        if stem in exact_names:
            candidates.append((100, f))
        elif normalize(book) in stem and (normalize(f"{ch}장") in stem or normalize(f"{ch2}장") in stem):
            candidates.append((50, f))
    if not candidates:
        raise FileNotFoundError(f"성경 PPT를 찾지 못했습니다: {book} {chapter}장 under {bible_root}")
    candidates.sort(key=lambda item: (-item[0], len(str(item[1])), str(item[1])))
    return candidates[0][1]


def strip_verse_number(text: str, verse: int) -> str:
    t = text.replace("\u00a0", " ")
    t = re.sub(r"\s+", " ", t).strip()
    t = re.sub(rf"^\s*{verse}\s*", "", t).strip()
    return t


def extract_verse_text(chapter_pptx: Path, verse: int) -> str:
    prs = Presentation(str(chapter_pptx))
    preferred_indices = [verse - 1] if 0 <= verse - 1 < len(prs.slides) else []
    indices = preferred_indices + [i for i in range(len(prs.slides)) if i not in preferred_indices]
    pattern = re.compile(rf"^\s*{verse}\s+(.+)", re.S)
    fallback: list[str] = []
    for idx in indices:
        slide = prs.slides[idx]
        for sh in slide.shapes:
            if not hasattr(sh, "text"):
                continue
            txt = sh.text.strip()
            if not txt:
                continue
            cleaned = re.sub(r"\s+", " ", txt.replace("\u00a0", " ")).strip()
            m = pattern.match(cleaned)
            if m and len(m.group(1).strip()) >= 2:
                return strip_verse_number(cleaned, verse)
            if cleaned.startswith(str(verse)) and len(cleaned) > len(str(verse)) + 2:
                fallback.append(cleaned)
    if fallback:
        fallback.sort(key=len, reverse=True)
        return strip_verse_number(fallback[0], verse)
    raise ValueError(f"{chapter_pptx.name}에서 {verse}절 텍스트를 찾지 못했습니다")


def chapter_verse_count(chapter_pptx: Path) -> int:
    return len(Presentation(str(chapter_pptx)).slides)


def expand_refs(ref: RefRange, chapter_files: dict[tuple[str, int], Path]) -> list[VerseRef]:
    refs: list[VerseRef] = []
    seen: set[tuple[str, int, int]] = set()
    for book, seg_sc, seg_sv, seg_ec, seg_ev in ref.iter_book_segments():
        for ch in range(seg_sc, seg_ec + 1):
            start = seg_sv if ch == seg_sc else 1
            end = seg_ev if ch == seg_ec else chapter_verse_count(chapter_files[(book, ch)])
            for verse in range(start, end + 1):
                key = (book, ch, verse)
                if key in seen:
                    continue
                seen.add(key)
                refs.append(VerseRef(book, ch, verse))
    return refs


def get_section_indices(pptx_path: Path, section_name: str) -> list[int]:
    prs = Presentation(str(pptx_path))
    sld_id_list = prs.slides._sldIdLst
    id_to_idx = {elem.get("id"): i for i, elem in enumerate(sld_id_list)}
    section_lst = prs._element.find(".//p14:sectionLst", {"p14": P14_NS})
    if section_lst is None:
        raise ValueError(f"섹션 목록이 없습니다: {pptx_path}")
    matches = []
    for section in section_lst.findall("p14:section", {"p14": P14_NS}):
        if section.get("name") != section_name:
            continue
        ids = [s.get("id") for s in section.findall(".//p14:sldId", {"p14": P14_NS})]
        indices = sorted(id_to_idx[sid] for sid in ids if sid in id_to_idx)
        matches.append(indices)
    if not matches:
        raise ValueError(f"섹션을 찾지 못했습니다: {section_name}")
    if len(matches) > 1:
        raise ValueError(f"동일한 이름의 섹션이 여러 개입니다: {section_name}")
    return matches[0]


def remove_shape(shape) -> None:
    el = shape._element
    el.getparent().remove(el)


def set_shape_name(shape, name: str) -> None:
    c_nv_pr = shape._element.find(f".//{{{PML_NS}}}cNvPr")
    if c_nv_pr is not None:
        c_nv_pr.set("name", name)


def emu(inches: float) -> int:
    return int(round(inches * EMU_PER_INCH))


def first_text_run_pr(shape):
    """Return a deep copy of the first existing run properties, if present."""
    tx_body = shape._element.find(f"{{{PML_NS}}}txBody")
    if tx_body is None:
        return None
    rpr = tx_body.find(f".//{{{A_NS}}}r/{{{A_NS}}}rPr")
    if rpr is not None:
        return copy.deepcopy(rpr)
    end_rpr = tx_body.find(f".//{{{A_NS}}}endParaRPr")
    if end_rpr is not None:
        return copy.deepcopy(end_rpr)
    return None


def first_text_p_pr(shape):
    """Return a deep copy of the first paragraph properties, if present."""
    tx_body = shape._element.find(f"{{{PML_NS}}}txBody")
    if tx_body is None:
        return None
    ppr = tx_body.find(f".//{{{A_NS}}}p/{{{A_NS}}}pPr")
    return copy.deepcopy(ppr) if ppr is not None else None


def font_size_from_rpr(rpr, fallback: float) -> float:
    if rpr is not None and rpr.get("sz"):
        try:
            return int(rpr.get("sz")) / 100.0
        except ValueError:
            pass
    return fallback


def set_rpr_font_size(rpr, size_pt: float | None) -> None:
    if rpr is not None and size_pt is not None:
        rpr.set("sz", str(int(round(size_pt * 100))))


def shape_text_height_pt(shape, *, default_t_ins: int = 45720, default_b_ins: int = 45720) -> float:
    """Return the approximate usable text box height in points."""
    body_pr = shape._element.find(f".//{{{A_NS}}}bodyPr")
    t_ins = default_t_ins
    b_ins = default_b_ins
    if body_pr is not None:
        if body_pr.get("tIns") is not None:
            t_ins = int(body_pr.get("tIns"))
        if body_pr.get("bIns") is not None:
            b_ins = int(body_pr.get("bIns"))
    usable_emu = max(0, int(shape.height) - t_ins - b_ins)
    return usable_emu / EMU_PER_INCH * 72.0


def shape_text_width_in(shape, *, default_l_ins: int = 91440, default_r_ins: int = 91440) -> float:
    """Return the approximate usable text box width in inches."""
    body_pr = shape._element.find(f".//{{{A_NS}}}bodyPr")
    l_ins = default_l_ins
    r_ins = default_r_ins
    if body_pr is not None:
        if body_pr.get("lIns") is not None:
            l_ins = int(body_pr.get("lIns"))
        if body_pr.get("rIns") is not None:
            r_ins = int(body_pr.get("rIns"))
    usable_emu = max(0, int(shape.width) - l_ins - r_ins)
    return usable_emu / EMU_PER_INCH


def shape_line_height_factor(shape, *, fallback: float = 1.3) -> float:
    """Return the template paragraph line-spacing factor, e.g. 150% -> 1.5."""
    ppr = shape._element.find(f".//{{{A_NS}}}p/{{{A_NS}}}pPr")
    ln_spc = ppr.find(f"{{{A_NS}}}lnSpc") if ppr is not None else None
    spc_pct = ln_spc.find(f"{{{A_NS}}}spcPct") if ln_spc is not None else None
    if spc_pct is not None and spc_pct.get("val"):
        try:
            # DrawingML stores 100% as 100000.
            return max(1.0, int(spc_pct.get("val")) / 100000.0)
        except ValueError:
            pass
    return fallback


def estimate_wrapped_line_count(text: str, font_size: float, width_in: float) -> int:
    if width_in <= 0:
        return max(1, len(text.splitlines()) or 1)
    lines = 0
    for paragraph in text.splitlines() or [text]:
        current = 0.0
        paragraph_lines = 1
        for char in paragraph:
            w = char_width_in(char, font_size)
            if current > 0 and current + w > width_in:
                paragraph_lines += 1
                current = 0.0 if char.isspace() else w
            else:
                current += w
        lines += max(1, paragraph_lines)
    return max(1, lines)


def wrap_text_by_words(text: str, font_size: float, width_in: float) -> str:
    """Greedy word wrapping that never splits a whitespace-delimited word."""
    if width_in <= 0:
        return text
    lines: list[str] = []
    for paragraph in text.splitlines() or [text]:
        words = paragraph.split()
        if not words:
            lines.append("")
            continue
        current = words[0]
        for word in words[1:]:
            candidate = f"{current} {word}"
            if sum(char_width_in(c, font_size) for c in candidate) <= width_in:
                current = candidate
            else:
                lines.append(current)
                current = word
        lines.append(current)
    return "\n".join(lines)


def fit_font_size_to_textbox_height(
    shape,
    text: str,
    *,
    default_size: float = 32.0,
    min_size: float = 16.0,
) -> float:
    """
    Default to 32pt and shrink only when wrapped text would exceed the
    textbox's vertical space. Width is used only to estimate wrapping lines.
    """
    available_pt = shape_text_height_pt(shape)
    if available_pt <= 0:
        return default_size
    width_in = shape_text_width_in(shape)
    line_height_factor = 1.2
    size = default_size
    while size > min_size:
        lines = estimate_wrapped_line_count(text, size, width_in)
        if lines * size * line_height_factor <= available_pt:
            return size
        size -= 0.5
    return min_size


def fit_body_text_to_textbox(
    shape,
    text: str,
    *,
    default_size: float = 32.0,
    min_size: float = 16.0,
    height_comfort_ratio: float = 0.80,
) -> tuple[str, float]:
    """
    Return manually wrapped body text and the largest comfortable font size.

    32pt remains the default for normal verses. For long verses, do not wait
    until the text barely overflows the textbox; shrink slightly once the text
    would consume the visual safe area. The textbox itself can extend into the
    grass/footer region, so use only about 80% of its nominal height and a
    conservative line-height estimate.
    """
    available_pt = shape_text_height_pt(shape)
    # Use a conservative usable width so rendered Korean text stays inside the
    # original text box and does not collide with the side decorations.
    width_in = shape_text_width_in(shape) * 0.86
    if available_pt <= 0:
        return wrap_text_by_words(text, default_size, width_in), default_size
    comfortable_pt = available_pt * height_comfort_ratio
    line_height_factor = shape_line_height_factor(shape)
    size = default_size
    while size > min_size:
        wrapped = wrap_text_by_words(text, size, width_in)
        lines = max(1, len(wrapped.splitlines()))
        if lines * size * line_height_factor <= comfortable_pt:
            return wrapped, size
        size -= 0.5
    return wrap_text_by_words(text, min_size, width_in), min_size


def ensure_verse_number_double_space(text: str) -> str:
    """Keep two literal spaces after the leading verse number on body slides."""
    return re.sub(r"^(\d+)\s+", r"\1  ", text, count=1)


def set_single_run_text_preserving_style(shape, text: str, *, size_pt: float | None = None) -> None:
    """
    Replace text while preserving the template textbox's existing paragraph and run style.

    This avoids hard-coding fonts. The font family, color, bold/italic, spacing,
    paragraph alignment, bullets/auto-numbering, and text-frame margins come from
    the textbox that already exists in the PPT template.
    """
    tx_body = shape._element.find(f"{{{PML_NS}}}txBody")
    if tx_body is None:
        raise ValueError(f"텍스트 본문이 없는 도형입니다: {getattr(shape, 'name', '<unknown>')}")

    first_p = tx_body.find(f"{{{A_NS}}}p")
    ppr = copy.deepcopy(first_p.find(f"{{{A_NS}}}pPr")) if first_p is not None and first_p.find(f"{{{A_NS}}}pPr") is not None else None
    rpr = None
    end_para_rpr = None
    if first_p is not None:
        first_r = first_p.find(f"{{{A_NS}}}r")
        if first_r is not None:
            rpr = first_r.find(f"{{{A_NS}}}rPr")
        if rpr is None:
            rpr = first_p.find(f"{{{A_NS}}}endParaRPr")
        end_para_rpr = first_p.find(f"{{{A_NS}}}endParaRPr")
    rpr = copy.deepcopy(rpr) if rpr is not None else etree.Element(f"{{{A_NS}}}rPr")
    end_para_rpr = copy.deepcopy(end_para_rpr) if end_para_rpr is not None else None
    set_rpr_font_size(rpr, size_pt)
    if end_para_rpr is not None:
        set_rpr_font_size(end_para_rpr, size_pt)

    for old_p in list(tx_body.findall(f"{{{A_NS}}}p")):
        tx_body.remove(old_p)

    new_p = etree.SubElement(tx_body, f"{{{A_NS}}}p")
    if ppr is not None:
        new_p.append(ppr)
    for i, line in enumerate(text.split("\n")):
        if i > 0:
            br = etree.SubElement(new_p, f"{{{A_NS}}}br")
            br.append(copy.deepcopy(rpr))
        new_r = etree.SubElement(new_p, f"{{{A_NS}}}r")
        new_r.append(copy.deepcopy(rpr))
        new_t = etree.SubElement(new_r, f"{{{A_NS}}}t")
        if line.startswith(" ") or line.endswith(" ") or "  " in line:
            new_t.set(f"{{{XML_NS}}}space", "preserve")
        new_t.text = line
    if end_para_rpr is not None:
        new_p.append(end_para_rpr)


def iter_shapes_recursive(shapes) -> Iterable:
    for sh in shapes:
        yield sh
        if sh.shape_type == MSO_SHAPE_TYPE.GROUP:
            yield from iter_shapes_recursive(sh.shapes)


def update_title_reference(slide, reference_text: str) -> None:
    for sh in iter_shapes_recursive(slide.shapes):
        if sh.name == "성경구절" and getattr(sh, "has_text_frame", False):
            # Multi-book references can be longer than the original template
            # textbox. Keep the existing font family/color/alignment, but widen
            # the textbox and disable wrapping so the reference stays on one
            # line as requested.
            if hasattr(sh, "left") and hasattr(sh, "width"):
                sh.left = emu(0.8)
                sh.width = emu(8.4)
            tx_body = sh._element.find(f"{{{PML_NS}}}txBody")
            body_pr = tx_body.find(f"{{{A_NS}}}bodyPr") if tx_body is not None else None
            if body_pr is not None:
                body_pr.set("wrap", "none")
            base_size = font_size_from_rpr(first_text_run_pr(sh), 35.0)
            visible = len(reference_text.replace(" ", ""))
            size = base_size
            if visible > 34:
                size = max(base_size - 7, 26)
            elif visible > 28:
                size = max(base_size - 5, 28)
            elif visible > 22:
                size = max(base_size - 3, 30)
            set_single_run_text_preserving_style(sh, reference_text, size_pt=size)
            return
    raise ValueError("제목 슬라이드에서 '성경구절' 텍스트 상자를 찾지 못했습니다")


def shape_text_from_el(shape_el) -> str:
    return "".join(t.text or "" for t in shape_el.findall(f".//{{{A_NS}}}t"))


def own_c_nv_pr(shape_el):
    if shape_el.tag == f"{{{PML_NS}}}grpSp":
        return shape_el.find(f"./{{{PML_NS}}}nvGrpSpPr/{{{PML_NS}}}cNvPr")
    return shape_el.find(f"./{{{PML_NS}}}nvSpPr/{{{PML_NS}}}cNvPr")


def shape_name_from_el(shape_el) -> str:
    c_nv_pr = own_c_nv_pr(shape_el)
    return c_nv_pr.get("name") if c_nv_pr is not None and c_nv_pr.get("name") else ""


def own_shape_top_in(shape_el) -> float | None:
    if shape_el.tag != f"{{{PML_NS}}}sp":
        return None
    off = shape_el.find(f"./{{{PML_NS}}}spPr/{{{A_NS}}}xfrm/{{{A_NS}}}off")
    if off is None or off.get("y") is None:
        return None
    try:
        return int(off.get("y")) / EMU_PER_INCH
    except ValueError:
        return None


def is_title_text_shape_el(shape_el) -> bool:
    if shape_el.tag != f"{{{PML_NS}}}sp":
        return False
    if shape_name_from_el(shape_el) == "성경구절":
        return False
    txt = shape_text_from_el(shape_el).strip()
    if not txt or txt in {"“", "”"}:
        return False
    y_in = own_shape_top_in(shape_el)
    return y_in is not None and 3.0 <= y_in <= 4.9


def find_group_element_by_name(root_el, name: str):
    for group_el in root_el.findall(f".//{{{PML_NS}}}grpSp"):
        c_nv_pr = own_c_nv_pr(group_el)
        if c_nv_pr is not None and c_nv_pr.get("name") == name:
            return group_el
    return None


def find_top_level_group_element(slide, name: str):
    sp_tree = slide.shapes._spTree
    for child in sp_tree:
        if child.tag == f"{{{PML_NS}}}grpSp" and shape_name_from_el(child) == name:
            return child
    return None


def find_title_template_shape_element(slide):
    title_group_el = find_top_level_group_element(slide, "본문 타이틀")
    search_roots = [title_group_el] if title_group_el is not None else []
    search_roots.append(slide.shapes._spTree)
    for root_el in search_roots:
        for shape_el in root_el.findall(f".//{{{PML_NS}}}sp"):
            if is_title_text_shape_el(shape_el):
                return shape_el
    return None


def find_quote_group_element(slide):
    title_group_el = find_top_level_group_element(slide, "본문 타이틀")
    if title_group_el is not None:
        nested_quote = find_group_element_by_name(title_group_el, "따옴표")
        if nested_quote is not None:
            return nested_quote
    return find_top_level_group_element(slide, "따옴표")


def first_text_run_pr_from_el(shape_el):
    rpr = shape_el.find(f".//{{{A_NS}}}r/{{{A_NS}}}rPr")
    if rpr is not None:
        return copy.deepcopy(rpr)
    end_rpr = shape_el.find(f".//{{{A_NS}}}endParaRPr")
    if end_rpr is not None:
        return copy.deepcopy(end_rpr)
    return None


def remove_old_title_block(slide) -> None:
    sp_tree = slide.shapes._spTree
    for child in list(sp_tree):
        if child.tag == f"{{{PML_NS}}}grpSp" and shape_name_from_el(child) in {"본문 타이틀", "따옴표"}:
            sp_tree.remove(child)
        elif is_title_text_shape_el(child):
            sp_tree.remove(child)


def max_shape_id(slide) -> int:
    ids = []
    for c_nv_pr in slide._element.xpath('.//p:cNvPr'):
        val = c_nv_pr.get('id')
        if val and val.isdigit():
            ids.append(int(val))
    return max(ids, default=1)


def renumber_shape_tree(shape_el, start_id: int) -> int:
    next_id = start_id
    for c_nv_pr in shape_el.findall(f".//{{{PML_NS}}}cNvPr"):
        c_nv_pr.set("id", str(next_id))
        next_id += 1
    return next_id


def make_empty_title_group(shape_id: int):
    group_el = etree.Element(f"{{{PML_NS}}}grpSp")
    nv_grp = etree.SubElement(group_el, f"{{{PML_NS}}}nvGrpSpPr")
    c_nv_pr = etree.SubElement(nv_grp, f"{{{PML_NS}}}cNvPr")
    c_nv_pr.set("id", str(shape_id))
    c_nv_pr.set("name", "본문 타이틀")
    etree.SubElement(nv_grp, f"{{{PML_NS}}}cNvGrpSpPr")
    etree.SubElement(nv_grp, f"{{{PML_NS}}}nvPr")
    grp_sp_pr = etree.SubElement(group_el, f"{{{PML_NS}}}grpSpPr")
    xfrm = etree.SubElement(grp_sp_pr, f"{{{A_NS}}}xfrm")
    etree.SubElement(xfrm, f"{{{A_NS}}}off", x=str(emu(0.3925)), y=str(emu(3.0913)))
    etree.SubElement(xfrm, f"{{{A_NS}}}ext", cx=str(emu(9.25)), cy=str(emu(1.92)))
    etree.SubElement(xfrm, f"{{{A_NS}}}chOff", x=str(emu(0.3925)), y=str(emu(3.0913)))
    etree.SubElement(xfrm, f"{{{A_NS}}}chExt", cx=str(emu(9.25)), cy=str(emu(1.92)))
    return group_el


def prepare_title_group(template_group_el, shape_id: int):
    if template_group_el is not None:
        group_el = copy.deepcopy(template_group_el)
        for child in list(group_el):
            if child.tag not in {f"{{{PML_NS}}}nvGrpSpPr", f"{{{PML_NS}}}grpSpPr"}:
                group_el.remove(child)
        c_nv_pr = own_c_nv_pr(group_el)
        if c_nv_pr is not None:
            c_nv_pr.set("id", str(shape_id))
            c_nv_pr.set("name", "본문 타이틀")
        return group_el
    return make_empty_title_group(shape_id)


def shape_id_name_maps(root_el) -> tuple[dict[str, str], dict[str, str]]:
    id_to_name: dict[str, str] = {}
    name_to_id: dict[str, str] = {}
    for c_nv_pr in root_el.findall(f".//{{{PML_NS}}}cNvPr"):
        sid = c_nv_pr.get("id")
        name = c_nv_pr.get("name")
        if not sid or not name:
            continue
        id_to_name[sid] = name
        name_to_id.setdefault(name, sid)
    return id_to_name, name_to_id


def copy_timing_by_shape_names(src_slide_el, dst_slide_el) -> None:
    """
    Copy the PowerPoint-authored animation timing from a template title slide.

    The user's PowerPoint-authored timing is the source of truth. Generated
    slides get new shape IDs, so every animation target is remapped by stable
    shape names such as 그룹 49, 성경구절, and 본문 타이틀.
    """
    src_timing = src_slide_el.find(f"{{{PML_NS}}}timing")
    if src_timing is None:
        raise ValueError("애니메이션 템플릿 제목 슬라이드에 timing XML이 없습니다")

    src_id_to_name, _ = shape_id_name_maps(src_slide_el)
    _, dst_name_to_id = shape_id_name_maps(dst_slide_el)
    timing = copy.deepcopy(src_timing)

    remapped = 0
    unresolved: list[str] = []
    for elem in timing.findall(f".//{{{PML_NS}}}spTgt") + timing.findall(f".//{{{PML_NS}}}bldP"):
        old_id = elem.get("spid")
        if not old_id:
            continue
        name = src_id_to_name.get(old_id)
        if name and name in dst_name_to_id:
            elem.set("spid", dst_name_to_id[name])
            remapped += 1
        else:
            unresolved.append(name or f"id={old_id}")

    if unresolved:
        raise ValueError(f"템플릿 애니메이션 대상 도형을 새 슬라이드에 연결하지 못했습니다: {', '.join(unresolved)}")
    if remapped == 0:
        raise ValueError("템플릿 애니메이션에서 복사할 도형 대상(spid)을 찾지 못했습니다")

    for old_timing in list(dst_slide_el.findall(f"{{{PML_NS}}}timing")):
        dst_slide_el.remove(old_timing)
    ext_lst = dst_slide_el.find(f"{{{PML_NS}}}extLst")
    if ext_lst is not None:
        dst_slide_el.insert(list(dst_slide_el).index(ext_lst), timing)
    else:
        dst_slide_el.append(timing)


def title_animation_template_elements(template_path: Path, section_name: str):
    if not template_path.exists():
        raise FileNotFoundError(f"애니메이션 템플릿 없음: {template_path}")
    indices = get_section_indices(template_path, section_name)
    if not indices:
        raise ValueError(f"애니메이션 템플릿에서 섹션을 찾지 못했습니다: {section_name}")
    prs = Presentation(str(template_path))
    first_el = copy.deepcopy(prs.slides[indices[0]]._element)
    last_el = copy.deepcopy(prs.slides[indices[-1]]._element)
    missing = []
    if first_el.find(f"{{{PML_NS}}}timing") is None:
        missing.append("첫 제목 슬라이드")
    if last_el.find(f"{{{PML_NS}}}timing") is None:
        missing.append("마지막 제목 슬라이드")
    if missing:
        raise ValueError(f"애니메이션 템플릿에 timing XML이 없습니다: {', '.join(missing)}")
    return first_el, last_el


def set_shape_position_and_text(shape_el, *, shape_id: int, name: str, x: int, y: int, cx: int, cy: int,
                                text: str, size_pt: float | None) -> None:
    c_nv_pr = shape_el.find(f".//{{{PML_NS}}}cNvPr")
    if c_nv_pr is not None:
        c_nv_pr.set("id", str(shape_id))
        c_nv_pr.set("name", name)

    xfrm = shape_el.find(f".//{{{A_NS}}}xfrm")
    if xfrm is None:
        sp_pr = shape_el.find(f"{{{PML_NS}}}spPr")
        if sp_pr is None:
            sp_pr = etree.SubElement(shape_el, f"{{{PML_NS}}}spPr")
        xfrm = etree.SubElement(sp_pr, f"{{{A_NS}}}xfrm")
    off = xfrm.find(f"{{{A_NS}}}off")
    if off is None:
        off = etree.SubElement(xfrm, f"{{{A_NS}}}off")
    ext = xfrm.find(f"{{{A_NS}}}ext")
    if ext is None:
        ext = etree.SubElement(xfrm, f"{{{A_NS}}}ext")
    off.set("x", str(x))
    off.set("y", str(y))
    ext.set("cx", str(cx))
    ext.set("cy", str(cy))

    texts = shape_el.findall(f".//{{{A_NS}}}t")
    if texts:
        texts[0].text = text
        for t in texts[1:]:
            t.text = ""
    else:
        tx_body = shape_el.find(f"{{{PML_NS}}}txBody")
        if tx_body is None:
            tx_body = etree.SubElement(shape_el, f"{{{PML_NS}}}txBody")
            etree.SubElement(tx_body, f"{{{A_NS}}}bodyPr")
            etree.SubElement(tx_body, f"{{{A_NS}}}lstStyle")
        p = etree.SubElement(tx_body, f"{{{A_NS}}}p")
        r = etree.SubElement(p, f"{{{A_NS}}}r")
        etree.SubElement(r, f"{{{A_NS}}}rPr")
        t = etree.SubElement(r, f"{{{A_NS}}}t")
        t.text = text

    for rpr in shape_el.findall(f".//{{{A_NS}}}rPr") + shape_el.findall(f".//{{{A_NS}}}endParaRPr"):
        set_rpr_font_size(rpr, size_pt)


def char_width_in(char: str, font_size: float) -> float:
    if char.isspace():
        return font_size / 135.0
    if re.match(r"[.,:;!?\-~'\"‘’“”()\[\]]", char):
        return font_size / 155.0
    if re.match(r"[A-Za-z0-9]", char):
        return font_size / 125.0
    return font_size / 90.0


def split_title_lines(title: str, font_size: float, max_width: float) -> list[str]:
    words = title.strip().split()
    if not words:
        return [""]
    lines: list[str] = []
    current = ""
    def width(s: str) -> float:
        return sum(char_width_in(c, font_size) for c in s)
    for w in words:
        cand = w if not current else current + " " + w
        if current and width(cand) > max_width:
            lines.append(current)
            current = w
        else:
            current = cand
    if current:
        lines.append(current)
    if len(lines) <= 2:
        return lines
    # If spaces create too many lines, rebalance by visible character count into 2 lines.
    raw = title.strip()
    mid = len(raw) // 2
    split_at = max(raw.rfind(" ", 0, mid + 1), 0)
    if split_at <= 0:
        split_at = mid
    return [raw[:split_at].strip(), raw[split_at:].strip()]


def choose_title_font(title: str, base_size: float) -> float:
    visible = len(title.replace(" ", ""))
    if visible <= 15:
        return base_size
    if visible <= 19:
        return max(base_size - 4, 28)
    if visible <= 24:
        return max(base_size - 8, 26)
    return max(base_size - 12, 24)


TITLE_PARTICLES = ("으로", "로", "에게", "께", "에서", "부터", "까지", "보다", "처럼", "과", "와", "을", "를", "이", "가", "은", "는", "의", "에")
TITLE_PUNCT_RE = re.compile(r"[.,:;!?\-~'\"‘’“”()\[\]]")


def title_word_stem(word: str) -> str:
    stripped = TITLE_PUNCT_RE.sub("", word)
    for particle in sorted(TITLE_PARTICLES, key=len, reverse=True):
        if stripped.endswith(particle) and len(stripped) > len(particle) + 1:
            return stripped[: -len(particle)]
    return stripped


def title_char_style(line: str, char_index: int, font_size: float) -> tuple[float, float]:
    """
    Return (size, y_offset_inches) for one title character.

    The layout stays deterministic but gives important Korean content words a
    little more weight. Particles and punctuation become slightly smaller.
    """
    if line[char_index].isspace():
        return font_size, 0.0

    # Find the whitespace-delimited word containing this character.
    start = char_index
    while start > 0 and not line[start - 1].isspace():
        start -= 1
    end = char_index + 1
    while end < len(line) and not line[end].isspace():
        end += 1
    word = line[start:end]
    local_idx = char_index - start
    stem = title_word_stem(word)
    char = line[char_index]

    if TITLE_PUNCT_RE.fullmatch(char):
        return max(font_size * 0.82, 22.0), 0.08

    stem_len = len(stem)
    if stem_len >= 2 and local_idx < stem_len:
        # Core word characters: larger and slightly lifted.
        factor = 1.14 if stem_len >= 3 else 1.08
        wave = [-0.04, -0.075, -0.045, -0.02][local_idx % 4]
        return font_size * factor, wave

    # Particles/endings: smaller and slightly lowered.
    return max(font_size * 0.88, 22.0), 0.065


def styled_line_metrics(line: str, font_size: float) -> list[tuple[str, float, float, float]]:
    """Return [(char, size_pt, y_offset_in, width_in), ...]."""
    metrics = []
    for i, char in enumerate(line):
        size, y_offset = title_char_style(line, i, font_size)
        metrics.append((char, size, y_offset, char_width_in(char, size)))
    return metrics


def fit_styled_line(line: str, font_size: float, max_width: float) -> tuple[list[tuple[str, float, float, float]], float]:
    metrics = styled_line_metrics(line, font_size)
    width = sum(item[3] for item in metrics)
    if width <= max_width or width <= 0:
        return metrics, width
    scale = max(max_width / width, 0.72)
    fitted = []
    for char, size, y_offset, _ in metrics:
        new_size = max(size * scale, 22.0)
        fitted.append((char, new_size, y_offset * scale, char_width_in(char, new_size)))
    return fitted, sum(item[3] for item in fitted)


def add_title_chars(slide, title: str, template_group_el, quote_group_el, template_char_el) -> None:
    if template_char_el is None:
        raise ValueError("제목 글자 템플릿 텍스트박스를 찾지 못했습니다")
    template_el = copy.deepcopy(template_char_el)
    base_rpr = first_text_run_pr_from_el(template_char_el)
    # A title slide can be regenerated from a PPT that already contains
    # generated per-character title boxes. If we use those generated sizes as
    # the next base, the 1.14x core-character styling compounds on every run.
    # Keep the design base stable so repeated runs are idempotent.
    base_size = min(font_size_from_rpr(base_rpr, TITLE_DESIGN_BASE_SIZE_PT), TITLE_DESIGN_BASE_SIZE_PT)
    font_size = choose_title_font(title, base_size)
    max_width = 8.8
    lines = split_title_lines(title, font_size, max_width)
    if len(lines) == 1:
        y_positions = [3.55]
    else:
        y_positions = [3.27, 4.05]

    next_id = max_shape_id(slide) + 1
    title_group_id = next_id
    next_id += 1
    title_group_el = prepare_title_group(template_group_el, title_group_id)
    if quote_group_el is not None:
        quote_copy = copy.deepcopy(quote_group_el)
        next_id = renumber_shape_tree(quote_copy, next_id)
        title_group_el.append(quote_copy)

    for line, y in zip(lines, y_positions):
        metrics, width = fit_styled_line(line, font_size, max_width)
        x = (10.0 - width) / 2.0
        for char, char_size, y_offset, w in metrics:
            if char.isspace():
                x += w
                continue
            new_el = copy.deepcopy(template_el)
            # Keep a visually stable centerline while allowing larger/smaller
            # characters to breathe up and down.
            char_y = y + y_offset - ((char_size - font_size) / 72.0) * 0.32
            set_shape_position_and_text(
                new_el,
                shape_id=next_id,
                name=f"설교제목글자 {next_id}",
                x=emu(x - 0.05),
                y=emu(char_y),
                cx=emu(max(0.45, w + 0.18)),
                cy=emu(max(0.72, char_size / 54.0)),
                text=char,
                size_pt=char_size,
            )
            title_group_el.append(new_el)
            next_id += 1
            x += w

    slide.shapes._spTree.insert_element_before(title_group_el, 'p:extLst')


def update_title_slide(slide, reference_text: str, title: str) -> None:
    update_title_reference(slide, reference_text)
    template_group_el = find_top_level_group_element(slide, "본문 타이틀")
    quote_group_el = find_quote_group_element(slide)
    template_char_el = find_title_template_shape_element(slide)
    remove_old_title_block(slide)
    add_title_chars(slide, title, template_group_el, quote_group_el, template_char_el)


def disable_body_auto_numbering(shape) -> None:
    """Remove automatic bullet/numbering from the body paragraph."""
    tx_body = shape._element.find(f".//{{{PML_NS}}}txBody")
    if tx_body is None:
        return
    first_p = tx_body.find(f"{{{A_NS}}}p")
    if first_p is None:
        return
    ppr = first_p.find(f"{{{A_NS}}}pPr")
    if ppr is None:
        ppr = etree.SubElement(first_p, f"{{{A_NS}}}pPr")
    for attr in ("marL", "indent"):
        if attr in ppr.attrib:
            del ppr.attrib[attr]
    for child in list(ppr):
        if child.tag.endswith(("}buAutoNum", "}buChar", "}buBlip", "}buNone", "}tabLst")):
            ppr.remove(child)
    etree.SubElement(ppr, f"{{{A_NS}}}buNone")


def update_verse_slide(slide, ref: VerseRef, verse_text: str) -> None:
    header = f"{ref.book} {ref.chapter}장"
    header_done = False
    body_done = False
    for sh in list(slide.shapes):
        if getattr(sh, "name", "") in UNUSED_VERSE_NAMES:
            remove_shape(sh)
    for sh in slide.shapes:
        if not getattr(sh, "has_text_frame", False):
            continue
        if sh.name in VERSE_HEADER_NAMES:
            set_shape_name(sh, "TextBox 1")
            set_single_run_text_preserving_style(sh, header)
            header_done = True
        elif sh.name in VERSE_BODY_NAMES:
            set_shape_name(sh, "TextBox 2")
            body_text = f"{ref.verse}  {verse_text}"
            wrapped_text, size = fit_body_text_to_textbox(sh, body_text, default_size=32.0, min_size=16.0)
            wrapped_text = ensure_verse_number_double_space(wrapped_text)
            tx_body = sh._element.find(f".//{{{PML_NS}}}txBody")
            body_pr = tx_body.find(f"{{{A_NS}}}bodyPr") if tx_body is not None else None
            if body_pr is not None:
                # Use our own word-boundary line breaks instead of PowerPoint's
                # automatic wrapping, which can split Korean words awkwardly.
                body_pr.set("wrap", "none")
            set_single_run_text_preserving_style(sh, wrapped_text, size_pt=size)
            disable_body_auto_numbering(sh)
            body_done = True
    if not header_done:
        raise ValueError("본문 슬라이드에서 왼쪽 상단 헤더(TextBox 1 또는 TextBox 8)를 찾지 못했습니다")
    if not body_done:
        raise ValueError("본문 슬라이드에서 본문 텍스트 상자(TextBox 2 또는 TextBox 15)를 찾지 못했습니다")

def capture_sections_by_indices(pptx_path: Path) -> list[dict]:
    with zipfile.ZipFile(pptx_path) as z:
        xml = z.read("ppt/presentation.xml")
    root = etree.fromstring(xml)
    slide_ids = root.xpath("./p:sldIdLst/p:sldId", namespaces=NS)
    id_to_idx = {e.get("id"): i for i, e in enumerate(slide_ids)}
    sections = []
    for sec in root.xpath(".//p14:sectionLst/p14:section", namespaces=NS):
        ids = [e.get("id") for e in sec.xpath("./p14:sldIdLst/p14:sldId", namespaces=NS)]
        indices = [id_to_idx[sid] for sid in ids if sid in id_to_idx]
        sections.append({"name": sec.get("name"), "id": sec.get("id"), "indices": indices})
    return sections


def restore_sections_by_indices(pptx_path: Path, sections: list[dict]) -> None:
    with zipfile.ZipFile(pptx_path, "r") as zin:
        files = {item.filename: zin.read(item) for item in zin.infolist()}
    root = etree.fromstring(files["ppt/presentation.xml"])
    slide_ids = root.xpath("./p:sldIdLst/p:sldId", namespaces=NS)

    # Remove existing p14 section extension, if any.
    for ext in root.xpath("./p:extLst/p:ext[p14:sectionLst]", namespaces=NS):
        ext.getparent().remove(ext)
    ext_lst = root.find(f"{{{PML_NS}}}extLst")
    if ext_lst is None:
        ext_lst = etree.SubElement(root, f"{{{PML_NS}}}extLst")
    ext = etree.SubElement(ext_lst, f"{{{PML_NS}}}ext")
    ext.set("uri", P14_EXT_URI)
    section_lst = etree.SubElement(ext, f"{{{P14_NS}}}sectionLst")
    for sec_data in sections:
        sec = etree.SubElement(section_lst, f"{{{P14_NS}}}section")
        sec.set("name", sec_data["name"] or "")
        sec.set("id", sec_data.get("id") or "{" + str(uuid.uuid4()).upper() + "}")
        sld_id_lst = etree.SubElement(sec, f"{{{P14_NS}}}sldIdLst")
        for idx in sec_data["indices"]:
            if 0 <= idx < len(slide_ids):
                sld = etree.SubElement(sld_id_lst, f"{{{P14_NS}}}sldId")
                sld.set("id", slide_ids[idx].get("id"))

    files["ppt/presentation.xml"] = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)
    tmp = pptx_path.with_suffix(pptx_path.suffix + ".sections")
    with zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as zout:
        for name, data in files.items():
            zout.writestr(name, data)
    os.replace(tmp, pptx_path)


def normalize_with_libreoffice(pptx_path: Path, tmp_dir: Path) -> None:
    sections = capture_sections_by_indices(pptx_path)
    out_dir = tmp_dir / "normalized"
    out_dir.mkdir(parents=True, exist_ok=True)
    result = subprocess.run(
        ["libreoffice", "--headless", "--convert-to", "pptx", "--outdir", str(out_dir), str(pptx_path)],
        text=True,
        capture_output=True,
    )
    normalized = out_dir / pptx_path.name
    if result.returncode != 0 or not normalized.exists():
        raise RuntimeError(f"LibreOffice 정규화 실패\nSTDOUT:{result.stdout}\nSTDERR:{result.stderr}")
    shutil.copy2(normalized, pptx_path)
    restore_sections_by_indices(pptx_path, sections)


def apply_title_slide_animations_to_pptx(pptx_path: Path, section_name: str, animation_template_path: Path) -> None:
    """Ensure first/last sermon title slides keep the template-authored animations after normalization."""
    sections = capture_sections_by_indices(pptx_path)
    indices = get_section_indices(pptx_path, section_name)
    if not indices:
        return
    first_anim_el, last_anim_el = title_animation_template_elements(animation_template_path, section_name)
    prs = Presentation(str(pptx_path))
    copy_timing_by_shape_names(first_anim_el, prs.slides[indices[0]]._element)
    copy_timing_by_shape_names(last_anim_el, prs.slides[indices[-1]]._element)
    tmp = pptx_path.with_suffix(pptx_path.suffix + ".anim")
    prs.save(str(tmp))
    os.replace(tmp, pptx_path)
    restore_sections_by_indices(pptx_path, sections)


def create_generated_sermon_pptx(target: Path, section_name: str, reference: RefRange, title: str,
                                 verse_texts: list[tuple[VerseRef, str]], out_pptx: Path,
                                 section_replacer_path: Path, animation_template_path: Path) -> None:
    helper = load_section_replacer(section_replacer_path)
    indices = get_section_indices(target, section_name)
    if len(indices) < 2:
        raise ValueError(f"설교 섹션에는 최소 제목/본문 템플릿 슬라이드가 필요합니다: {indices}")
    title_template_idx = indices[0]
    verse_template_idx = indices[1] if len(indices) >= 2 else indices[0]
    closing_title_template_idx = indices[-1]

    src = Presentation(str(target))
    dest = Presentation()
    dest.slide_width = src.slide_width
    dest.slide_height = src.slide_height

    # python-pptx new presentations have no slides by default.
    helper.copy_slide(src, title_template_idx, dest)
    for _ in verse_texts:
        helper.copy_slide(src, verse_template_idx, dest)
    helper.copy_slide(src, closing_title_template_idx, dest)

    ref_text = display_reference(reference)
    update_title_slide(dest.slides[0], ref_text, title)
    for i, (vref, text) in enumerate(verse_texts, start=1):
        update_verse_slide(dest.slides[i], vref, text)
    update_title_slide(dest.slides[len(dest.slides) - 1], ref_text, title)
    first_anim_el, last_anim_el = title_animation_template_elements(animation_template_path, section_name)
    copy_timing_by_shape_names(first_anim_el, dest.slides[0]._element)
    copy_timing_by_shape_names(last_anim_el, dest.slides[len(dest.slides) - 1]._element)

    helper.save_deduplicated(dest, str(out_pptx))


def replace_section_with_source(target: Path, section_name: str, source: Path, section_replacer_path: Path) -> None:
    cmd = [sys.executable, str(section_replacer_path), "--replace", "--target", str(target), "--section", section_name, "--source", str(source)]
    result = subprocess.run(cmd, text=True, capture_output=True)
    if result.returncode != 0:
        raise RuntimeError(f"섹션 교체 실패\nCMD: {' '.join(cmd)}\nSTDOUT:{result.stdout}\nSTDERR:{result.stderr}")
    print(result.stdout.strip())


def verify_pptx(path: Path) -> tuple[int, int]:
    with zipfile.ZipFile(path) as z:
        bad = z.testzip()
    if bad:
        raise RuntimeError(f"PPTX zip 오류: {bad}")
    prs = Presentation(str(path))
    sections = capture_sections_by_indices(path)
    return len(prs.slides), len(sections)


def main() -> int:
    ap = argparse.ArgumentParser(description="설교 섹션을 잔디/꽃 제목 템플릿 + 구절별 본문 슬라이드로 교체")
    ap.add_argument("--target", required=True, help="작업할 예배 PPTX")
    ap.add_argument("--reference", required=True, help="본문 범위 예: '사사기 4:1-10'")
    ap.add_argument("--title", required=True, help="설교 제목")
    ap.add_argument("--section", default="설교", help="교체할 섹션명 (기본: 설교)")
    ap.add_argument("--bible-root", default=DEFAULT_BIBLE_ROOT, help="개역개정성경ppt 루트 폴더")
    ap.add_argument("--section-replacer", default=DEFAULT_SECTION_REPLACER, help="church-worship-ppt 섹션 교체 스크립트")
    ap.add_argument("--animation-template", default=DEFAULT_ANIMATION_TEMPLATE, help="제목 슬라이드 애니메이션을 복사할 template.pptx")
    ap.add_argument("--no-normalize", action="store_true", help="LibreOffice 정규화와 구역 복원을 건너뜀")
    ap.add_argument("--keep-temp", action="store_true", help="임시 생성 PPTX를 삭제하지 않음")
    args = ap.parse_args()

    target = as_path(args.target)
    bible_root = as_path(args.bible_root)
    section_replacer = as_path(args.section_replacer)
    animation_template = as_path(args.animation_template)
    if not target.exists():
        raise FileNotFoundError(f"대상 PPTX 없음: {target}")

    ref_range = parse_reference(args.reference)
    with tempfile.TemporaryDirectory(prefix="sermon_grass_") as td:
        tmp_dir = Path(td)
        chapter_files_raw: dict[tuple[str, int], Path] = {}
        chapter_files_pptx: dict[tuple[str, int], Path] = {}
        needed_chapters = []
        seen_chapters: set[tuple[str, int]] = set()
        for book, sc, _, ec, _ in ref_range.iter_book_segments():
            for ch in range(sc, ec + 1):
                key = (book, ch)
                if key not in seen_chapters:
                    seen_chapters.add(key)
                    needed_chapters.append(key)
        for book, ch in needed_chapters:
            raw = find_chapter_ppt(bible_root, book, ch)
            chapter_files_raw[(book, ch)] = raw
            chapter_files_pptx[(book, ch)] = convert_to_pptx(raw, tmp_dir)

        refs = expand_refs(ref_range, chapter_files_pptx)
        verse_texts = [
            (vref, extract_verse_text(chapter_files_pptx[(vref.book, vref.chapter)], vref.verse))
            for vref in refs
        ]

        generated = tmp_dir / "generated_sermon.pptx"
        create_generated_sermon_pptx(target, args.section, ref_range, args.title, verse_texts, generated, section_replacer, animation_template)
        replace_section_with_source(target, args.section, generated, section_replacer)
        if not args.no_normalize:
            normalize_with_libreoffice(target, tmp_dir)
        apply_title_slide_animations_to_pptx(target, args.section, animation_template)
        slides, section_count = verify_pptx(target)

        if args.keep_temp:
            keep_dir = target.parent / f".{target.stem}_sermon_grass_temp"
            if keep_dir.exists():
                shutil.rmtree(keep_dir)
            shutil.copytree(tmp_dir, keep_dir)
            print(f"TEMP_DIR: {keep_dir}")

    print("완료")
    print(f"TARGET: {target}")
    print(f"SECTION: {args.section}")
    print(f"REFERENCE: {display_reference(ref_range)}")
    print(f"TITLE: {args.title}")
    print(f"VERSE_SLIDES: {len(verse_texts)}")
    for (book, ch), raw in chapter_files_raw.items():
        print(f"BIBLE_SOURCE: {book} {ch}장 -> {raw}")
    print(f"TOTAL_SLIDES: {slides}")
    print(f"SECTIONS: {section_count}")
    return 0


if __name__ == "__main__":
    try:
        raise SystemExit(main())
    except Exception as e:
        print(f"ERROR: {e}", file=sys.stderr)
        raise SystemExit(1)
